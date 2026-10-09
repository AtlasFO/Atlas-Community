"""Tests for tools/tabular.py — typed xlsx/csv/jsonl evidence access."""
import json

import pytest

from tools.tabular import (table_grep, table_pivot, table_query, table_schema)


def _fn(tool):
    return getattr(tool, "fn", tool)


def test_a_row_with_a_very_large_field_is_readable(tmp_path):
    """A Prefetch row lists every file the program loaded; a field of a
    few hundred kilobytes is ordinary parsed evidence. The csv module's
    default limit refused it from every table tool at once."""
    p = tmp_path / "prefetch.csv"
    loaded = ",".join(f"\\VOLUME\\WINDOWS\\SYSTEM32\\DLL{i:06d}.DLL" for i in range(9000))
    assert len(loaded) > 131072
    p.write_text("ExecutableName,RunCount,FilesLoaded\n"
                 f"EVIL.EXE,3,\"{loaded}\"\n", encoding="utf-8")
    schema = _fn(table_schema)(str(p))
    assert schema["success"] is True, schema
    r = _fn(table_query)(str(p), where=["ExecutableName=EVIL.EXE"])
    assert r["success"] is True, r
    assert r["matched_rows"] == 1


def test_zero_match_with_a_doubled_backslash_names_the_cause(tmp_path):
    """A pattern that arrived double-escaped matches nothing; the zero stays
    honest and the result says why, so the retry is the right one."""
    p = tmp_path / "t.csv"
    p.write_text("ip,n\n10.0.0.1,1\n", encoding="utf-8")
    doubled = _fn(table_query)(path=str(p), where=["ip~^\\\\d+"], limit=5)
    assert doubled["success"] and doubled["matched_rows"] == 0
    assert "doubled backslash" in doubled["hint"]
    single = _fn(table_query)(path=str(p), where=["ip~^\\d+"], limit=5)
    assert single["matched_rows"] == 1 and "doubled" not in single["hint"]


@pytest.fixture()
def siem_xlsx(tmp_path):
    """Two-sheet workbook shaped like a SIEM export."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alerts"
    ws.append(["timestamp", "src_ip", "dst_ip", "severity", "rule"])
    ws.append(["2031-02-04 10:00:00", "10.0.0.5", "203.0.113.9", "high",
               "beacon detected"])
    ws.append(["2031-02-04 10:05:00", "10.0.0.5", "203.0.113.9", "high",
               "beacon detected"])
    ws.append(["2031-02-04 11:00:00", "10.0.0.7", "198.51.100.2", "low",
               "dns anomaly"])
    ws2 = wb.create_sheet("Assets")
    ws2.append(["host", "owner"])
    ws2.append(["PC01", "j.doe"])
    p = tmp_path / "siem_export.xlsx"
    wb.save(p)
    return str(p)


@pytest.fixture()
def events_csv(tmp_path):
    p = tmp_path / "events.csv"
    p.write_text(
        "user,host,event_id\n"
        "j.doe,PC01,4624\n"
        "j.doe,PC01,4624\n"
        "svc.backup01,FILESRV02,4672\n")
    return str(p)


@pytest.fixture()
def cloud_jsonl(tmp_path):
    p = tmp_path / "cloudtrail.jsonl"
    lines = [
        {"eventName": "ConsoleLogin", "sourceIPAddress": "203.0.113.9",
         "userIdentity": "admin"},
        {"eventName": "RunInstances", "sourceIPAddress": "10.0.0.5",
         "userIdentity": "dev"},
    ]
    p.write_text("\n".join(json.dumps(x) for x in lines))
    return str(p)


class TestSchema:
    def test_xlsx_sheets_columns_rows(self, siem_xlsx):
        r = _fn(table_schema)(siem_xlsx)
        assert r["success"] is True
        assert set(r["sheets"]) == {"Alerts", "Assets"}
        assert r["sheets"]["Alerts"]["columns"] == [
            "timestamp", "src_ip", "dst_ip", "severity", "rule"]
        assert r["sheets"]["Alerts"]["rows"] == 3
        assert r["sheets"]["Alerts"]["examples"]["src_ip"] == "10.0.0.5"

    def test_csv(self, events_csv):
        r = _fn(table_schema)(events_csv)
        assert r["sheets"]["-"]["columns"] == ["user", "host", "event_id"]
        assert r["sheets"]["-"]["rows"] == 3

    def test_jsonl(self, cloud_jsonl):
        r = _fn(table_schema)(cloud_jsonl)
        assert "eventName" in r["sheets"]["-"]["columns"]

    def test_missing_file(self):
        assert _fn(table_schema)("/nope.xlsx")["success"] is False

    def test_unsupported_extension(self, tmp_path):
        p = tmp_path / "image.e01"
        p.write_bytes(b"\x00")
        r = _fn(table_schema)(str(p))
        assert r["success"] is False
        assert r.get("gate") == "wrong_input_kind"
        assert "tabular export" in r["error"].lower()


class TestElasticExportQuirks:
    """Shapes of SIEM exports:
    Kibana xlsx puts each record's whole CSV line into one cell, and leaves
    the human-readable timestamp unquoted so its comma shifts every field."""

    @pytest.fixture()
    def kibana_xlsx(self, tmp_path):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "in"
        ws.append(['@timestamp,"user.name","event.action","host.name"'])
        ws.append(['Feb 4, 2031 @ 12:00:01.000,svc_reader,session-created,node-02'])
        ws.append(['Feb 4, 2031 @ 12:05:00.000,"admin_svc",logout,node-02'])
        p = tmp_path / "kibana.xlsx"
        wb.save(p)
        return str(p)

    def test_embedded_csv_header_and_rows_align(self, kibana_xlsx):
        r = _fn(table_schema)(kibana_xlsx)
        assert r["sheets"]["in"]["columns"] == [
            "@timestamp", "user.name", "event.action", "host.name"]
        q = _fn(table_query)(kibana_xlsx, where=["user.name=svc_reader"])
        assert q["matched_rows"] == 1
        row = q["rows"][0]
        assert row["event.action"] == "session-created"
        assert row["host.name"] == "node-02"
        assert row["@timestamp"].startswith("Feb 4, 2031")


class TestQuery:
    def test_equals_and_columns(self, siem_xlsx):
        r = _fn(table_query)(siem_xlsx, sheet="Alerts",
                             where=["severity=HIGH"],
                             columns=["src_ip", "rule"])
        assert r["matched_rows"] == 2
        assert r["rows"][0] == {"src_ip": "10.0.0.5", "rule": "beacon detected"}

    def test_regex_and_negation(self, siem_xlsx):
        r = _fn(table_query)(siem_xlsx, sheet="Alerts",
                             where=["rule~beacon|dns", "severity!=high"])
        assert r["matched_rows"] == 1
        assert r["rows"][0]["severity"] == "low"

    def test_numeric_comparison(self, events_csv):
        r = _fn(table_query)(events_csv, where=["event_id>4624"])
        assert r["matched_rows"] == 1
        assert r["rows"][0]["event_id"] == "4672"

    def test_sort_and_limit(self, siem_xlsx):
        r = _fn(table_query)(siem_xlsx, sheet="Alerts", sort_by="timestamp",
                             descending=True, limit=1)
        assert r["returned_rows"] == 1
        assert r["rows"][0]["timestamp"].startswith("2031-02-04 11")
        assert r["matched_rows"] == 3  # caller sees there is more than the window

    def test_bad_predicate_is_an_error(self, events_csv):
        r = _fn(table_query)(events_csv, where=["no-operator-here"])
        assert r["success"] is False
        assert "unparseable predicate" in r["error"]

    def test_unknown_columns_refused(self, events_csv):
        # Truly unbound names (not in synonym groups) exhaust after auto-peek.
        r = _fn(table_query)(
            events_csv,
            columns=["user", "not_a_real_column"],
            where=["not_a_real_column=user01"],
        )
        assert r["success"] is False
        assert r.get("gate") == "schema_columns_absent"
        assert "not_a_real_column" in r.get("unknown_columns", [])
        assert r.get("schema_columns")

    def test_unambiguous_synonym_rewritten(self, events_csv):
        # username → user (singular identity column); eid → event_id
        r = _fn(table_query)(
            events_csv,
            columns=["username", "hostname", "eid"],
            where=["username=j.doe"],
        )
        assert r["success"] is True
        assert r["matched_rows"] == 2
        assert r.get("columns_rewritten", {}).get("username") == "user"
        assert r["columns_rewritten"].get("hostname") == "host"
        assert r["columns_rewritten"].get("eid") == "event_id"
        assert "user" in r["rows"][0]

    def test_ambiguous_synonym_refused(self, tmp_path):
        # Two user-like headers → refuse free-form `user` (no guessing).
        p = tmp_path / "dual_user.csv"
        p.write_text(
            "SubjectUserName,TargetUserName,host\n"
            "j.doe,user01,PC01\n"
        )
        r = _fn(table_query)(str(p), where=["user=j.doe"])
        assert r["success"] is False
        assert r.get("gate") == "schema_columns_absent"
        assert "user" in r.get("unknown_columns", [])
        assert r.get("schema_columns")

    def test_timestamp_prefers_receive_time_when_ambiguous(self, tmp_path):
        # Firewall-export dual time columns: Timestamp → Receive Time (not refuse).
        p = tmp_path / "vpn.csv"
        p.write_text(
            "#Domain,Receive Time,Generate Time,Source User\n"
            "x,2031-02-04 12:00:00,2031-02-04 12:00:01,admin\n"
        )
        r = _fn(table_query)(
            str(p), columns=["Timestamp", "Source User"],
        )
        assert r["success"] is True
        assert r.get("columns_rewritten", {}).get("Timestamp") == "Receive Time"
        assert "Source User" in r["rows"][0] or "source user" in {
            k.lower() for k in r["rows"][0]
        }

    def test_source_ip_synonym_with_dual_ip_headers(self, siem_xlsx):
        # src_ip + dst_ip present: sourceip binds; bare destinationip binds.
        r = _fn(table_query)(
            siem_xlsx, sheet="Alerts", where=["sourceip=10.0.0.5"],
        )
        assert r["success"] is True
        assert r["matched_rows"] == 2
        assert r.get("columns_rewritten", {}).get("sourceip") == "src_ip"
        r2 = _fn(table_query)(
            siem_xlsx, sheet="Alerts", where=["destinationip=203.0.113.9"],
        )
        assert r2["success"] is True
        assert r2["matched_rows"] == 2

    def test_schema_cache_persists_under_case(self, tmp_path):
        case = tmp_path / "case"
        (case / "evidence").mkdir(parents=True)
        (case / "analysis").mkdir()
        (case / ".atlas").mkdir()
        p = case / "evidence" / "events.csv"
        p.write_text("user,host,event_id\nj.doe,PC01,4624\n")
        schema = _fn(table_schema)(str(p))
        assert schema["success"] is True
        cache = case / ".atlas" / "tabular_schema_cache.json"
        assert cache.is_file()
        # Query after schema uses cached headers (rewrite still works).
        r = _fn(table_query)(str(p), where=["username=j.doe"])
        assert r["success"] is True
        assert r["matched_rows"] == 1

    def test_casefold_column_canonicalized(self, events_csv):
        # headers are user,host,event_id — accept USER=
        r = _fn(table_query)(events_csv, where=["USER=j.doe"])
        assert r["success"] is True
        assert r["matched_rows"] == 2

    def test_backtick_column_names_accepted(self, events_csv):
        # Models emit SQL-style `user`='j.doe' — must not unknown_columns
        r = _fn(table_query)(events_csv, where=["`user`='j.doe'"])
        assert r["success"] is True
        assert r["matched_rows"] == 2
        r2 = _fn(table_query)(
            events_csv, columns=["`user`", "`host`"], where=["user=j.doe"],
        )
        assert r2["success"] is True
        assert "user" in r2["rows"][0]

    def test_sql_or_tree_refused(self, events_csv):
        r = _fn(table_query)(
            events_csv,
            where=["`user`='j.doe' OR `user`='admin'"],
        )
        assert r["success"] is False
        assert r.get("gate") == "unsupported_where_dialect"
        assert r.get("rewrite_hint")
        assert "two table.table_query" in r["rewrite_hint"]

    def test_auto_peek_exhausts_invented_columns(self, tmp_path):
        """B1 residual: peek establishes schema SoT — invent → exhaust in-band."""
        case = tmp_path / "case"
        (case / "evidence").mkdir(parents=True)
        (case / "analysis").mkdir()
        (case / ".atlas").mkdir()
        p = case / "evidence" / "firewall.csv"
        p.write_text(
            "#Domain,Receive Time,Source User,Device\n"
            "corp,2031-02-04 12:00:00,admin,fw01\n"
        )
        bad = _fn(table_query)(
            str(p), where=["SourceIP=203.0.113.4", "SessionName=x"],
        )
        assert bad.get("gate") == "schema_columns_absent"
        assert bad.get("schema_columns")
        assert "Receive Time" in (bad.get("schema_columns") or [])
        # Valid query proceeds without a mandatory table_schema hop.
        ok = _fn(table_query)(
            str(p),
            columns=["Timestamp", "Source User", "Device", "Domain"],
            where=["Source User=admin"],
        )
        assert ok["success"] is True
        assert ok["matched_rows"] == 1
        assert ok.get("columns_rewritten", {}).get("Timestamp") == "Receive Time"
        assert ok.get("columns_rewritten", {}).get("Domain") == "#Domain"
        assert ok.get("schema_columns")
        invent = _fn(table_query)(
            str(p), where=["LogonType=3", "FailureReason=x"],
        )
        assert invent.get("gate") == "schema_columns_absent"
        assert invent.get("schema_columns")
        assert invent.get("closest_columns") is not None

    def test_edr_destination_synonyms_bind_to_remote(self, tmp_path):
        """I4: Destination* product names → Remote* when unique in file."""
        p = tmp_path / "edr_logons.csv"
        p.write_text(
            "TimeGenerated,RemoteDeviceName,RemoteIP,RemotePort,AccountName\n"
            "2031-02-04T12:00:00Z,host02,10.0.0.101,445,user01\n",
            encoding="utf-8",
        )
        r = _fn(table_query)(
            str(p),
            columns=[
                "DestinationDeviceName",
                "DestinationIPAddress",
                "DestinationPort",
            ],
            where=["DestinationIPAddress=10.0.0.101"],
        )
        assert r["success"] is True, r
        assert r.get("columns_rewritten", {}).get(
            "DestinationDeviceName"
        ) == "RemoteDeviceName"
        assert r.get("columns_rewritten", {}).get(
            "DestinationIPAddress"
        ) == "RemoteIP"
        assert r.get("columns_rewritten", {}).get(
            "DestinationPort"
        ) == "RemotePort"
        assert r["matched_rows"] == 1
        assert "RemoteDeviceName" in (r.get("schema_columns") or [])

    def test_device_synonym_binds_to_host_like_header(self, tmp_path):
        p = tmp_path / "edr.csv"
        p.write_text(
            "TimeGenerated,DeviceName,AccountName\n"
            "2031-02-04T12:00:00Z,host01,j.doe\n"
        )
        r = _fn(table_query)(
            str(p), columns=["Timestamp", "Device", "user"],
        )
        assert r["success"] is True
        assert r.get("columns_rewritten", {}).get("Device") == "DeviceName"
        assert r.get("columns_rewritten", {}).get("Timestamp") == "TimeGenerated"
        assert r.get("columns_rewritten", {}).get("user") == "AccountName"

    def test_non_numeric_compare_refused(self, events_csv):
        r = _fn(table_query)(
            events_csv,
            where=["event_id>='Feb 4, 2031'"],
        )
        assert r["success"] is False
        assert r.get("gate") == "non_numeric_compare"

    def test_valid_zero_still_success(self, events_csv):
        r = _fn(table_query)(events_csv, where=["user=nobody"])
        assert r["success"] is True
        assert r["matched_rows"] == 0
        assert r.get("valid_zero") is True

    def test_output_csv_spills_all_matches(self, siem_xlsx, tmp_path):
        out = tmp_path / "analysis" / "high.csv"
        r = _fn(table_query)(siem_xlsx, sheet="Alerts",
                             where=["severity=high"], limit=1,
                             output_csv=str(out))
        assert r["output_csv"] == str(out)
        content = out.read_text().splitlines()
        assert len(content) == 3  # header + 2 matches (uncapped)

    def test_output_csv_skips_zero_match_spill(self, events_csv, tmp_path):
        """Zero matches must not create a 2-byte \\r\\n headerless file."""
        out = tmp_path / "analysis" / "nobody.csv"
        r = _fn(table_query)(events_csv, where=["user=nobody"],
                             output_csv=str(out))
        assert r["success"] is True
        assert r["matched_rows"] == 0
        assert r.get("output_csv") in (None, "")
        assert not out.exists()

    def test_refuses_empty_tabular_spill(self, tmp_path):
        empty = tmp_path / "analysis" / "empty.csv"
        empty.parent.mkdir(parents=True)
        empty.write_bytes(b"\r\n")
        r = _fn(table_query)(str(empty), where=["user=x"])
        assert r["success"] is False
        assert r.get("gate") == "empty_tabular"

    def test_output_csv_refuses_evidence_path(self, siem_xlsx):
        with pytest.raises(ValueError):
            _fn(table_query)(siem_xlsx, sheet="Alerts",
                             output_csv="/mnt/host01/out.csv")


class TestPivot:
    def test_group_counts(self, siem_xlsx):
        r = _fn(table_pivot)(siem_xlsx, group_by=["src_ip", "dst_ip"],
                             sheet="Alerts")
        assert r["groups"][0] == {"src_ip": "10.0.0.5",
                                  "dst_ip": "203.0.113.9", "count": 2}
        assert r["distinct_groups"] == 2

    def test_pivot_with_where(self, events_csv):
        r = _fn(table_pivot)(events_csv, group_by=["user"],
                             where=["event_id=4624"])
        assert r["groups"] == [{"user": "j.doe", "count": 2}]


class TestGrep:
    def test_ioc_sweep_across_sheets(self, siem_xlsx):
        r = _fn(table_grep)(siem_xlsx, r"203\.0\.113\.9")
        assert r["hit_count"] == 2
        assert all("dst_ip" in h["_matched_columns"] for h in r["hits"])
        assert r["hits"][0]["_sheet"] == "Alerts"

    def test_jsonl_grep(self, cloud_jsonl):
        r = _fn(table_grep)(cloud_jsonl, "ConsoleLogin")
        assert r["hit_count"] == 1
        assert r["hits"][0]["userIdentity"] == "admin"

    def test_bad_regex(self, events_csv):
        r = _fn(table_grep)(events_csv, "[unclosed")
        assert r["success"] is False


class TestScanBound:
    """A scan reads every row unless an operator bounds it; a bounded scan
    says how far it read, and finding nothing before the bound is
    inconclusive, not a negative."""

    @pytest.fixture()
    def timeline(self, tmp_path):
        p = tmp_path / "CORP-DC01-timeline.csv"
        rows = [f"2031-01-0{1 + i % 9} 10:00:{i:02d},logon,user{i:02d},{i}" for i in range(12)]
        rows[10] = "2031-01-09 10:00:10,logon,jane.doe,10"
        p.write_text("time,event,user,n\n" + "\n".join(rows) + "\n", encoding="utf-8")
        return str(p)

    def test_the_whole_file_is_scanned_by_default(self, timeline):
        r = _fn(table_grep)(timeline, "jane\\.doe")
        assert r["hit_count"] == 1 and r["scan_capped"] is False
        assert r["scanned_rows"] == 12 and "inconclusive" not in r

    def test_a_bounded_scan_that_found_nothing_is_inconclusive(self, timeline, monkeypatch):
        import tools.tabular as tabular
        monkeypatch.setattr(tabular, "MAX_SCAN_ROWS", 5)
        r = _fn(table_grep)(timeline, "jane\\.doe")
        assert r["hit_count"] == 0 and r["scan_capped"] is True and r["scanned_rows"] == 5
        assert r["inconclusive"] is True and "first 5 rows" in r["note"]
        _fn(table_schema)(timeline)
        q = _fn(table_query)(timeline, where=["user=jane.doe"])
        assert q["matched_rows"] == 0 and q["valid_zero"] is False and q["inconclusive"] is True

    def test_a_query_holds_only_its_window_and_counts_every_match(self, timeline):
        _fn(table_schema)(timeline)
        r = _fn(table_query)(timeline, sort_by="n", descending=True, limit=3, offset=2)
        assert r["matched_rows"] == 12
        assert [row["n"] for row in r["rows"]] == ["9", "8", "7"]
        r = _fn(table_query)(timeline, limit=2, offset=1)
        assert r["matched_rows"] == 12 and [row["n"] for row in r["rows"]] == ["1", "2"]
